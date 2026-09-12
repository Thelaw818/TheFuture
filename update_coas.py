#!/usr/bin/env python3
import copy
import hashlib
import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
DB_PATH = ROOT / "coa-database.json"

def rec(id, compound, ctype, labeled, actual, purity, status, lab, date, report,
        key, batch, source, notes, crimp, cap, closure, identity="Confirmed"):
    variance = round(actual - labeled, 3)
    return [id, compound, ctype, labeled, actual, variance, variance / labeled,
            purity, status, lab, identity, date, report, key, batch, source,
            "Pass", "Unique product/strength/date/report record", notes, crimp, cap, closure]

records = [
    rec("COA-0056", "Tesamorelin", "Growth Hormone Support", 10, 12.18, 99.79,
        "Overfilled", "Pristin Biotech", "2026-09-01", "Not Listed", "Not Listed", "20260822",
        "assets/coa/documents/pristin-biotech/tesamorelin/Tesamorelin-10mg-2026-09-01-Pristin.pdf",
        "Identity detected by LCMS/MS; HPLC purity 99.79%; net peptide content 12.18 mg measured by nitrogen quantification using an elemental analyzer. White powder; received 2026-08-29 and analyzed 2026-09-01.",
        "Not Listed", "Not Listed", "Closure colors are not listed or visible in the report."),
    rec("COA-0057", "Retatrutide", "GLP-1 / Weight", 30, 37.37, 99.86,
        "Overfilled", "Pristin Biotech", "2026-09-01", "Not Listed", "Not Listed", "20260822",
        "assets/coa/documents/pristin-biotech/retatrutide/Retatrutide-30mg-2026-09-01-Pristin.pdf",
        "Identity detected by LCMS/MS; HPLC purity 99.86%; net peptide content 37.37 mg measured by nitrogen quantification using an elemental analyzer. White powder; received 2026-08-29 and analyzed 2026-09-01.",
        "Not Listed", "Not Listed", "Closure colors are not listed or visible in the report."),
    rec("COA-0058", "Retatrutide", "GLP-1 / Weight", 10, 10.80, 99.83,
        "Overfilled", "Pristin Biotech", "2026-09-01", "Not Listed", "Not Listed", "20260822",
        "assets/coa/documents/pristin-biotech/retatrutide/Retatrutide-10mg-2026-09-01-Pristin.pdf",
        "Identity detected by LCMS/MS; HPLC purity 99.83%; net peptide content 10.80 mg measured by nitrogen quantification using an elemental analyzer. White powder; received 2026-08-29 and analyzed 2026-09-01.",
        "Not Listed", "Not Listed", "Closure colors are not listed or visible in the report."),
    rec("COA-0059", "Glutathione", "Antioxidant / Cellular Support", 1500, 1478.47, 99.513,
        "On label", "Kovera Labs", "2026-08-18", "Redacted", "Redacted", "GLUTATI1500-2608-01-RED-SLV",
        "assets/coa/images/kovera-labs/glutathione/Glutathione-1500mg-2026-08-18-Kovera.jpg",
        "Identity confirmed as Glutathione by LC-MS. Endotoxin PASS; microbial sterility: No Growth; arsenic, cadmium, lead, and mercury Negative. Vial 1: 99.246% / 1464.59 mg; vial 2: 99.781% / 1492.36 mg.",
        "Silver", "Red", "Red flip-off cap with silver aluminum crimp, as listed in the COA."),
    rec("COA-0060", "Semax", "Nootropic / Research", 10, 9.65, 99.568,
        "On label", "Kovera Labs", "2026-08-18", "Redacted", "Redacted", "SMX10-2608-01-BLK-SLV",
        "assets/coa/images/kovera-labs/semax/Semax-10mg-2026-08-18-Kovera.jpg",
        "Identity confirmed as Semax by LC-MS. Endotoxin PASS; microbial sterility: No Growth; arsenic, cadmium, lead, and mercury Negative. Vial 1: 99.328% / 9.56 mg; vial 2: 99.808% / 9.74 mg.",
        "Silver", "Black", "Black flip-off cap with silver aluminum crimp, as listed in the COA."),
    rec("COA-0061", "Sermorelin", "Growth Hormone Support", 10, 10.41, 99.335,
        "On label", "Kovera Labs", "2026-08-18", "Redacted", "Redacted", "SERMOR10-2608-01-GRN-SLV",
        "assets/coa/images/kovera-labs/sermorelin/Sermorelin-10mg-2026-08-18-Kovera.jpg",
        "Identity confirmed as Sermorelin by LC-MS. Endotoxin PASS; microbial sterility: No Growth; arsenic, cadmium, lead, and mercury Negative. Vial 1: 99.266% / 10.31 mg; vial 2: 99.403% / 10.51 mg.",
        "Silver", "Green", "Green flip-off cap with silver aluminum crimp, as listed in the COA."),
    rec("COA-0062", "GHK-Cu", "Cosmetic / Recovery", 50, 52.48, 99.25,
        "On label", "SteriGenix", "2026-08-26", "RPT-2026-353561", "SGX-VP-2026-0812-U80", "CU260725A",
        "assets/coa/documents/sterigenix/ghk-cu/COA-RPT-2026-353561-SteriGenix.pdf",
        "Identity confirmed as GHK-Cu (Copper Peptide) by LC-MS. Blue lyophilized powder; molecular weight matched 402.92 g/mol. Endotoxin PASS; microbial sterility: No Growth; lead, cadmium, arsenic, and mercury Negative. Batch conformity was not tested.",
        "Silver", "Blue", "Blue flip-off cap with silver aluminum crimp, as listed in the COA."),
]

detail_extras = {
    "COA-0059": [["Identity", "Confirmed: Glutathione (LC-MS)"], ["Form", "Lyophilized Powder"], ["Endotoxin Safety Screen", "PASS"], ["Microbial Sterility Screen", "No Growth"], ["Heavy Metals", "Arsenic, Cadmium, Lead, Mercury: Negative"], ["Vial 1", "99.246% purity / 1464.59 mg"], ["Vial 2", "99.781% purity / 1492.36 mg"]],
    "COA-0060": [["Identity", "Confirmed: Semax (LC-MS)"], ["Form", "Lyophilized Powder"], ["Endotoxin Safety Screen", "PASS"], ["Microbial Sterility Screen", "No Growth"], ["Heavy Metals", "Arsenic, Cadmium, Lead, Mercury: Negative"], ["Vial 1", "99.328% purity / 9.56 mg"], ["Vial 2", "99.808% purity / 9.74 mg"]],
    "COA-0061": [["Identity", "Confirmed: Sermorelin (LC-MS)"], ["Form", "Lyophilized Powder"], ["Endotoxin Safety Screen", "PASS"], ["Microbial Sterility Screen", "No Growth"], ["Heavy Metals", "Arsenic, Cadmium, Lead, Mercury: Negative"], ["Vial 1", "99.266% purity / 10.31 mg"], ["Vial 2", "99.403% purity / 10.51 mg"]],
    "COA-0062": [["Identity", "Confirmed: GHK-Cu (Copper Peptide) (LC-MS)"], ["Form", "Lyophilized"], ["Appearance", "Blue Powder"], ["Molecular Formula", "C14H23CuN6O4+"], ["Molecular Weight", "402.92 g/mol — Matches"], ["CAS Number", "89030-95-5"], ["Endotoxin", "PASS"], ["Microbial Sterility Screen", "No Growth"], ["Heavy Metals", "Lead, Cadmium, Arsenic, Mercury: Negative"], ["Batch Conformity", "Not Tested"]],
}

def map_headers(headers, row):
    return {str(k): row[i] if i < len(row) else None for i, k in enumerate(headers)}

def iso_date(value):
    if isinstance(value, (int, float)):
        return datetime.fromordinal(datetime(1899, 12, 30).toordinal() + int(value)).date().isoformat()
    try:
        return datetime.fromisoformat(str(value)).date().isoformat()
    except ValueError:
        return ""

def update_dashboard(sheets, coa):
    body = coa[1:]
    compounds = sorted({r[1] for r in body if r[1]})
    labs = sorted({r[9] for r in body if r[9]})
    purities = [float(r[7]) for r in body if isinstance(r[7], (int, float))]
    latest = max((iso_date(r[11]) for r in body if r[11]), default="")
    counts = {"Overfilled": 0, "Underfilled": 0, "On Label/Review": 0}
    for r in body:
        s = str(r[8]).lower()
        if "over" in s: counts["Overfilled"] += 1
        elif "under" in s: counts["Underfilled"] += 1
        else: counts["On Label/Review"] += 1
    d = sheets.get("Dashboard", [])
    vals = {
        "Total COA Records": (len(body), "Synced from COA Database"),
        "Unique Peptides": (len(compounds), ", ".join(compounds)),
        "Average Purity %": (round(sum(purities)/len(purities), 3), "Populated purity results"),
        "Overfilled": (counts["Overfilled"], None), "Underfilled": (counts["Underfilled"], None),
        "On Label/Review": (counts["On Label/Review"], None),
        "Labs Used": (len(labs), ", ".join(labs)), "Latest COA Date": (latest, None)
    }
    for row in d:
        if row and row[0] in vals:
            row[1], note = vals[row[0]]
            if note is not None and len(row) > 2: row[2] = note

def rebuild_sheets(sheets):
    coa = sheets["COA Database"]
    h = coa[0]
    body = coa[1:]
    idx = {v:i for i,v in enumerate(h)}
    sheets["Batch History"] = copy.deepcopy(coa)
    cert_h = ["COA ID","Compound","Compound Type","Lab","Report/Task #","Verification Key","File/Link","Verified?","Labeled Qty (mg)","Actual Content (mg)","Variance (mg)","Variance %","Purity %","Status","Batch/Lot","Test Date","Pass/Fail","Notes","Crimp Color","Cap Color","Closure Notes"]
    sheets["Certificates"] = [cert_h] + [[r[idx["COA ID"]],r[idx["Compound"]],r[idx["Compound Type"]],r[idx["Lab"]],r[idx["Report/Task #"]],r[idx["Verification Key"]],r[idx["Source Type"]],r[idx["Lab Match"]],r[idx["Labeled Qty (mg)"]],r[idx["Actual Content (mg)"]],r[idx["Variance (mg)"]],r[idx["Variance %"]],r[idx["Purity %"]],r[idx["Status"]],r[idx["Batch/Lot"]],r[idx["Test Date"]],r[idx["Pass/Fail"]],r[idx["Notes"]],r[idx["Crimp Color"]],r[idx["Cap Color"]],r[idx["Closure Notes"]]] for r in body]
    inv_h = ["Item ID","Compound","Compound Type","Strength (mg)","Actual Content (mg)","Variance (mg)","Variance %","Purity %","Batch/Lot","Lab","Status","Last COA Date","COA ID","Report/Task #","Verification Key","Certificate Link","Cap Color","Crimp Color","Closure Notes","Notes"]
    sheets["Inventory"] = [inv_h] + [[r[idx["COA ID"]],r[idx["Compound"]],r[idx["Compound Type"]],r[idx["Labeled Qty (mg)"]],r[idx["Actual Content (mg)"]],r[idx["Variance (mg)"]],r[idx["Variance %"]],r[idx["Purity %"]],r[idx["Batch/Lot"]],r[idx["Lab"]],r[idx["Status"]],r[idx["Test Date"]],r[idx["COA ID"]],r[idx["Report/Task #"]],r[idx["Verification Key"]],r[idx["Source Type"]],r[idx["Cap Color"]],r[idx["Crimp Color"]],r[idx["Closure Notes"]],r[idx["Notes"]]] for r in body]
    purity_h = ["Rank","COA ID","Compound","Compound Type","Labeled Qty (mg)","Actual Content (mg)","Variance %","Purity %","Status","Lab","Test Date","Cap Color","Crimp Color","Notes"]
    ranked = sorted(body, key=lambda r: (-(float(r[idx["Purity %"]]) if isinstance(r[idx["Purity %"]],(int,float)) else -1), r[idx["COA ID"]]))
    sheets["Purity Rankings"] = [purity_h] + [[n,r[idx["COA ID"]],r[idx["Compound"]],r[idx["Compound Type"]],r[idx["Labeled Qty (mg)"]],r[idx["Actual Content (mg)"]],r[idx["Variance %"]],r[idx["Purity %"]],r[idx["Status"]],r[idx["Lab"]],r[idx["Test Date"]],r[idx["Cap Color"]],r[idx["Crimp Color"]],r[idx["Notes"]]] for n,r in enumerate(ranked,1)]
    status_h = ["COA ID","Compound","Compound Type","Labeled Qty (mg)","Actual Content (mg)","Variance (mg)","Variance %","Purity %","Status","Lab","Test Date","Report/Task #","Verification Key","Batch/Lot","Pass/Fail","Crimp Color","Cap Color","Closure Notes","Notes"]
    compounds = sorted({r[idx["Compound"]] for r in body}); labs=sorted({r[idx["Lab"]] for r in body}); pur=[r[idx["Purity %"]] for r in body if isinstance(r[idx["Purity %"]],(int,float))]
    metrics = [["COA Status Summary — Tested Peptides"]+[None]*18,["Metric","Value","Notes"]+[None]*16,
      ["Total COA/Test Records",len(body),"All records from COA Database"]+[None]*16,
      ["Unique Peptides",len(compounds),", ".join(compounds)]+[None]*16,
      ["Passed Records",sum(str(r[idx["Pass/Fail"]]).lower()=="pass" for r in body),"Pass/Fail column"]+[None]*16,
      ["Overfilled",sum("over" in str(r[idx["Status"]]).lower() for r in body),"Status column"]+[None]*16,
      ["Underfilled",sum("under" in str(r[idx["Status"]]).lower() for r in body),"Status column"]+[None]*16,
      ["On Label/Review",sum("over" not in str(r[idx["Status"]]).lower() and "under" not in str(r[idx["Status"]]).lower() for r in body),"Status column"]+[None]*16,
      ["Average Purity %",round(sum(pur)/len(pur),3),"Based on populated purity values"]+[None]*16,
      ["Unique Labs",len(labs),", ".join(labs)]+[None]*16,
      ["Last Test Date",max((iso_date(r[idx["Test Date"]]) for r in body if r[idx["Test Date"]]), default=""),"Most recent test date"]+[None]*16,
      [None]*19,status_h]
    status_rows = [[r[idx["COA ID"]],r[idx["Compound"]],r[idx["Compound Type"]],r[idx["Labeled Qty (mg)"]],r[idx["Actual Content (mg)"]],r[idx["Variance (mg)"]],r[idx["Variance %"]],r[idx["Purity %"]],r[idx["Status"]],r[idx["Lab"]],r[idx["Test Date"]],r[idx["Report/Task #"]],r[idx["Verification Key"]],r[idx["Batch/Lot"]],r[idx["Pass/Fail"]],r[idx["Crimp Color"]],r[idx["Cap Color"]],r[idx["Closure Notes"]],r[idx["Notes"]]] for r in body]
    sheets["COA Status"] = metrics + status_rows
    update_dashboard(sheets, coa)

def detail_sheet(row):
    h = ["COA ID","Compound","Compound Type","Labeled Qty (mg)","Actual Content (mg)","Variance (mg)","Variance %","Purity %","Status","Lab","Lab Match","Test Date","Report/Task #","Verification Key","Batch/Lot","Source Type","Pass/Fail","Duplicate Handling","Notes","Crimp Color","Cap Color","Closure Notes"]
    d = map_headers(h, row)
    pairs = [["Field","Value"],["COA ID",d["COA ID"]],["Product",d["Compound"]],["Compound Type",d["Compound Type"]],["Lab",d["Lab"]],["Test Date",d["Test Date"]],["Report/Task #",d["Report/Task #"]],["Verification Key",d["Verification Key"]],["Batch/Lot",d["Batch/Lot"]],["Labeled Qty (mg)",d["Labeled Qty (mg)"]],["Actual Content (mg)",d["Actual Content (mg)"]],["Variance (mg)",d["Variance (mg)"]],["Variance %",d["Variance %"]],["Purity %",d["Purity %"]],["Status",d["Status"]],["Cap Color",d["Cap Color"]],["Crimp Color",d["Crimp Color"]],["Closure Notes",d["Closure Notes"]],["Pass/Fail",d["Pass/Fail"]],["Notes",d["Notes"]],["Source",d["Source Type"]]]
    pairs[15:15] = detail_extras.get(d["COA ID"], [])
    return pairs

def update_json():
    data = json.loads(DB_PATH.read_text(encoding="utf-8"))
    sheets = data["sheets"]
    coa = sheets["COA Database"]
    existing = {r[0]:r for r in coa[1:]}
    for r in records:
        if r[0] not in existing: coa.append(r)
    # Link the two pre-existing Freedom Retatrutide records to their supplied images.
    image_updates = {"COA-0054":"assets/coa/images/freedom-diagnostics/retatrutide/Retatrutide-20mg-2026-08-29-Freedom.jpg", "COA-0055":"assets/coa/images/freedom-diagnostics/retatrutide/Retatrutide-40mg-2026-08-29-Freedom.jpg"}
    for r in coa[1:]:
        if r[0] in image_updates: r[15] = image_updates[r[0]]
    for r in records:
        sheets[f"COA_{r[0].replace('-','_')}"] = detail_sheet(r)
    for cid,path in image_updates.items():
        key=f"COA_{cid.replace('-','_')}"
        if key in sheets:
            found=False
            for pair in sheets[key]:
                if pair and pair[0] in ("Source Type","Image","PDF","Source"):
                    pair[1]=path; found=True
            if not found: sheets[key].append(["Image",path])
    rebuild_sheets(sheets)
    data["version"] = "2026-09-12-coa-0062-complete-sync"
    DB_PATH.write_text(json.dumps(data, indent=2, ensure_ascii=False)+"\n", encoding="utf-8")
    return data

def replace_sheet(ws, rows):
    style_row = 2 if ws.max_row >= 2 else 1
    styles = []
    for c in range(1, max(ws.max_column, max(map(len,rows)))+1):
        cell=ws.cell(style_row,c)
        styles.append((copy.copy(cell._style), copy.copy(cell.font), copy.copy(cell.fill), copy.copy(cell.border), copy.copy(cell.alignment), cell.number_format, copy.copy(cell.protection)))
    if ws.max_row: ws.delete_rows(1, ws.max_row)
    for rnum,row in enumerate(rows,1):
        for cnum,val in enumerate(row,1):
            cell=ws.cell(rnum,cnum,val)
            if rnum>1 and cnum<=len(styles):
                st=styles[cnum-1]; cell._style=copy.copy(st[0]); cell.number_format=st[5]
    ws.auto_filter.ref = f"A1:{ws.cell(1,len(rows[0])).column_letter}{len(rows)}" if rows and len(rows)>1 else None
    ws.freeze_panes = "A2"

def update_workbook(path, sheets):
    wb=load_workbook(path)
    for name in ["Dashboard","COA Database","Purity Rankings","Batch History","Inventory","Certificates","COA Status"]:
        if name in wb.sheetnames: replace_sheet(wb[name], sheets[name])
    for r in records:
        name=f"COA_{r[0].replace('-','_')}"
        if name in wb.sheetnames: del wb[name]
        ws=wb.create_sheet(name)
        replace_sheet(ws,sheets[name])
        ws.column_dimensions['A'].width=30; ws.column_dimensions['B'].width=105
    # Update the two existing detail tabs with image paths.
    for cid in ("COA-0054","COA-0055"):
        name=f"COA_{cid.replace('-','_')}"
        if name in wb.sheetnames:
            replace_sheet(wb[name],sheets[name])
    wb.save(path)

def update_index(sheets):
    p=ROOT/"index.html"; text=p.read_text(encoding="utf-8")
    payload=json.dumps(sheets,ensure_ascii=False,separators=(",",":"))
    text,n=re.subn(r"const EMBEDDED_WORKBOOK_DATA = \{.*?\};\nconst PRISTIN_COA_UPDATE", "const EMBEDDED_WORKBOOK_DATA = "+payload+";\nconst PRISTIN_COA_UPDATE", text, count=1, flags=re.S)
    if n!=1: raise RuntimeError("Embedded database marker not found")
    p.write_text(text,encoding="utf-8")

def update_gallery():
    p=ROOT/"coa-pictures.html"; text=p.read_text(encoding="utf-8")
    marker="const NEW_ITEMS = ["
    additions='''\n  {preview:"images/kovera-labs/glutathione/Glutathione-1500mg-2026-08-18-Kovera.jpg", pdf:null, label:"Glutathione 1500mg — COA-0059 — Red cap / Silver crimp"},\n  {preview:"images/kovera-labs/semax/Semax-10mg-2026-08-18-Kovera.jpg", pdf:null, label:"Semax 10mg — COA-0060 — Black cap / Silver crimp"},\n  {preview:"images/kovera-labs/sermorelin/Sermorelin-10mg-2026-08-18-Kovera.jpg", pdf:null, label:"Sermorelin 10mg — COA-0061 — Green cap / Silver crimp"},\n  {preview:"images/sterigenix/ghk-cu/COA-RPT-2026-353561-page-1.png", pdf:"documents/sterigenix/ghk-cu/COA-RPT-2026-353561-SteriGenix.pdf", label:"GHK-Cu 50mg — COA-0062 — Page 1"},\n  {preview:"images/sterigenix/ghk-cu/COA-RPT-2026-353561-page-2.png", pdf:"documents/sterigenix/ghk-cu/COA-RPT-2026-353561-SteriGenix.pdf", label:"GHK-Cu 50mg — COA-0062 — Page 2"},\n  {preview:"images/sterigenix/ghk-cu/COA-RPT-2026-353561-page-3.png", pdf:"documents/sterigenix/ghk-cu/COA-RPT-2026-353561-SteriGenix.pdf", label:"GHK-Cu 50mg — COA-0062 — Verification"},\n  {preview:"images/freedom-diagnostics/retatrutide/Retatrutide-20mg-2026-08-29-Freedom.jpg", pdf:null, label:"Retatrutide 20mg — COA-0054 — Clear Pink cap"},\n  {preview:"images/freedom-diagnostics/retatrutide/Retatrutide-40mg-2026-08-29-Freedom.jpg", pdf:null, label:"Retatrutide 40mg — COA-0055 — Clear Purple cap"},'''
    if "Glutathione 1500mg — COA-0059" not in text: text=text.replace(marker,marker+additions,1)
    p.write_text(text,encoding="utf-8")

def update_manifest():
    p=ROOT/"COA_FILE_MANIFEST.json"
    data=json.loads(p.read_text(encoding="utf-8"))
    existing={x.get("path") for x in data.get("files",[])}
    for f in sorted((ROOT/"assets/coa").rglob("*")):
        if not f.is_file(): continue
        rel=f.relative_to(ROOT).as_posix()
        if rel in existing: continue
        if any(x in rel for x in ("Glutathione-1500mg-2026-08-18", "Semax-10mg-2026-08-18", "Sermorelin-10mg-2026-08-18", "RPT-2026-353561", "Retatrutide-20mg-2026-08-29", "Retatrutide-40mg-2026-08-29")):
            lab="Kovera Labs" if "kovera-labs" in rel else "Freedom Diagnostics" if "freedom-diagnostics" in rel else "SteriGenix"
            product="Glutathione" if "Glutathione" in rel else "Semax" if "Semax" in rel else "Sermorelin" if "Sermorelin" in rel else "Retatrutide" if "Retatrutide" in rel else "GHK-Cu"
            data["files"].append({"path":rel,"category":"COA PDF" if f.suffix.lower()==".pdf" else "COA image","lab":lab,"product":product,"role":"document" if f.suffix.lower()==".pdf" else "image","sha256":hashlib.sha256(f.read_bytes()).hexdigest(),"bytes":f.stat().st_size})
    data["total_files"]=len(data["files"]); data["generated"]="2026-09-12T00:00:00Z"
    p.write_text(json.dumps(data,indent=2)+"\n",encoding="utf-8")

if __name__ == "__main__":
    data=update_json()
    for name in ("inventory.xlsx","inventory_updated.xlsx"):
        update_workbook(ROOT/name,data["sheets"])
    update_index(data["sheets"])
    update_gallery()
    update_manifest()
    print(f"Updated {len(data['sheets']['COA Database'])-1} COAs")
